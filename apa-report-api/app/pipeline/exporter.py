"""Render a completed report as .docx (python-docx) or .pdf (HTML via WeasyPrint), in APA 7 layout."""

from __future__ import annotations

import html
import io
import re

import markdown as md
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from app.schemas.output import Report

_INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|_[^_]+_)")


# ----------------------------------------------------------------------- docx
def _add_inline(paragraph, text: str) -> None:
    for part in _INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif (part.startswith("*") and part.endswith("*")) or (part.startswith("_") and part.endswith("_")):
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def _add_markdown(doc, content: str) -> None:
    for block in re.split(r"\n\s*\n", content.strip()):
        lines = [ln.rstrip() for ln in block.splitlines() if ln.strip()]
        if not lines:
            continue
        if all(re.match(r"^\s*([-*]|\d+\.)\s+", ln) for ln in lines):
            for ln in lines:
                style = "List Number" if re.match(r"^\s*\d+\.", ln) else "List Bullet"
                _add_inline(doc.add_paragraph(style=style), re.sub(r"^\s*([-*]|\d+\.)\s+", "", ln))
            continue
        heading = re.match(r"^(#{2,6})\s+(.*)$", lines[0])
        if heading and len(lines) == 1:
            p = doc.add_paragraph()
            run = p.add_run(heading.group(2).strip())
            run.bold = True
            if len(heading.group(1)) >= 4:
                run.italic = True
            continue
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Inches(0.5)
        _add_inline(p, " ".join(ln.strip() for ln in lines))


def render_docx(report: Report) -> bytes:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 2.0
    normal.paragraph_format.space_after = Pt(0)
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Inches(1)
        section.left_margin = section.right_margin = Inches(1)

    # Title page
    for _ in range(3):
        doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run(report.title).bold = True
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run(report.central_question).italic = True
    doc.add_page_break()

    if report.abstract:
        heading = doc.add_paragraph()
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        heading.add_run("Abstract").bold = True
        doc.add_paragraph(report.abstract)
        if report.keywords:
            kw = doc.add_paragraph()
            kw.paragraph_format.first_line_indent = Inches(0.5)
            kw.add_run("Keywords: ").italic = True
            kw.add_run(", ".join(report.keywords))
        doc.add_page_break()

    centered = doc.add_paragraph()
    centered.alignment = WD_ALIGN_PARAGRAPH.CENTER
    centered.add_run(report.title).bold = True
    for section in report.sections:
        heading = doc.add_paragraph()
        if section.heading.lower() != "introduction":
            heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
            heading.add_run(section.heading).bold = True
        _add_markdown(doc, section.content)

    if report.references:
        doc.add_page_break()
        refs_heading = doc.add_paragraph()
        refs_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        refs_heading.add_run("References").bold = True
        for ref in report.references:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.5)
            p.paragraph_format.first_line_indent = Inches(-0.5)
            _add_inline(p, ref.apa_citation)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# ------------------------------------------------------------------------ pdf
_PDF_CSS = """
@page { size: Letter; margin: 1in; @top-right { content: counter(page); font: 12pt 'Times New Roman', serif; } }
body { font-family: 'Times New Roman', Times, serif; font-size: 12pt; line-height: 2; }
h1, h2 { font-size: 12pt; font-weight: bold; text-align: center; margin: 0; }
h3 { font-size: 12pt; font-weight: bold; margin: 0; }
h4 { font-size: 12pt; font-weight: bold; font-style: italic; margin: 0; }
p { text-indent: 0.5in; margin: 0; }
.title-page { page-break-after: always; text-align: center; padding-top: 2in; }
.title-page p { text-indent: 0; }
.abstract { page-break-after: always; }
.abstract p.body { text-indent: 0; }
.references { page-break-before: always; }
.references p { padding-left: 0.5in; text-indent: -0.5in; }
"""


def report_html(report: Report) -> str:
    esc = html.escape
    parts = [
        "<!doctype html><html lang='en'><head><meta charset='utf-8'>",
        f"<title>{esc(report.title)}</title><style>{_PDF_CSS}</style></head><body>",
        f"<section class='title-page'><h1>{esc(report.title)}</h1><p><em>{esc(report.central_question)}</em></p></section>",
    ]
    if report.abstract:
        parts.append(f"<section class='abstract'><h2>Abstract</h2><p class='body'>{esc(report.abstract)}</p>")
        if report.keywords:
            parts.append(f"<p><em>Keywords:</em> {esc(', '.join(report.keywords))}</p>")
        parts.append("</section>")
    parts.append(f"<h1>{esc(report.title)}</h1>")
    for section in report.sections:
        heading_tag = "h3" if section.heading.lower() == "introduction" else "h2"
        if section.heading.lower() != "introduction":
            parts.append(f"<{heading_tag}>{esc(section.heading)}</{heading_tag}>")
        # Markdown from the pipeline; raw HTML is escaped first so it renders as text.
        parts.append(md.markdown(esc(section.content, quote=False), extensions=["sane_lists"]))
    if report.references:
        parts.append("<section class='references'><h2>References</h2>")
        for ref in report.references:
            parts.append(md.markdown(esc(ref.apa_citation, quote=False)))
        parts.append("</section>")
    parts.append("</body></html>")
    return "".join(parts)


def render_pdf(report: Report) -> bytes:
    from weasyprint import HTML  # imported lazily: loads native libraries

    return HTML(string=report_html(report)).write_pdf()
